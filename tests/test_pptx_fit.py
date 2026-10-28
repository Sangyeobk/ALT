from build_pptx import bullets_height, bullets_pages


def test_bullets_pages_keeps_emphasis_first_and_every_item_once():
    long = "가" * 400
    items = [(long, False), (long, False), ("선정 흐름", True), ("짧은 줄", False), ("분석 편수", True), (long, False)]
    pages = bullets_pages(items, (18.6, 13.8), (31.8, 13.7), 12)
    assert ("선정 흐름", True) in pages[0] and ("분석 편수", True) in pages[0]
    assert sorted(i for p in pages for i in p) == sorted(items)
    order = [items.index(i) for i in pages[0]]
    assert order == sorted(order)  # 장 안에서는 원래 순서
    assert all(bullets_height(p, 31.8, 12) <= 13.7 for p in pages[1:])


def test_bullets_pages_one_page_when_short():
    items = [("가나다", False), ("라마바", True)]
    assert bullets_pages(items, (18.6, 13.8), (31.8, 13.7), 12) == [items]


def test_sections_height_adds_one_line_per_heading():
    from build_pptx import PT_CM, sections_height
    items = [("짧은 줄", False), ("가" * 300, True)]
    head = 14 * PT_CM * 1.2 + 4 * PT_CM  # 머리말 (1)(2)(3) 한 줄
    assert abs(sections_height([("(1) 목적", items), ("(2) 결과", [])], 31.8, 14)
               - (2 * head + bullets_height(items, 31.8, 14))) < 1e-9


def test_deck_without_template_uses_plain_layout(tmp_path, monkeypatch):
    """양식 파일이 없어도 (공개 저장소에서 받은 사람) plain으로 표지·본문·마무리 순서로 만들어진다."""
    from pptx import Presentation
    import build_pptx
    from build_pptx import Deck, pptx_style
    config = build_pptx.load_config()
    missing = {**config, "output": {**config["output"], "pptx_template": "x",
                                    "pptx_templates": {**config["output"]["pptx_templates"], "x": {"file": "templates/없는_양식.pptx"}}}}
    monkeypatch.setattr(build_pptx, "load_config", lambda: missing)
    style = pptx_style()
    assert style["path"] is None and style["colors"] == config["output"]["pptx_templates"]["plain"]["colors"]
    deck = Deck(style)
    deck.set_cover(["제목 한 줄"], "머리글")
    deck.slide("Background", "[배경] 요지")
    deck.finish(tmp_path / "out.pptx")
    slides = list(Presentation(str(tmp_path / "out.pptx")).slides)
    assert len(slides) == 3
    texts = [" ".join(sh.text_frame.text for sh in s.shapes if sh.has_text_frame) for s in slides]
    assert "제목 한 줄" in texts[0] and "Background" in texts[1] and "Thank you" in texts[2]


def test_template_01_cover_section_body_closing(tmp_path):
    """양식 01 (사용자 결정 2026-10-21): 표지 → 장 구분 → 본문(테두리 복사) → 마무리(출처 표시 유지, 예시 연락처 지움),
    양식의 나머지 장은 남지 않음, 글꼴은 맑은 고딕. 양식 파일은 저장소에 없으므로 없으면 건너뜀."""
    import pytest
    from pptx import Presentation
    from build_pptx import FONT, Deck, pptx_style
    style = pptx_style("01")
    if style["path"] is None:
        pytest.skip("templates/ 양식 01 없음")
    deck = Deck(style)
    deck.set_cover(["LITERATURE SURVEY", "해양구조물 파랑하중 — 연구 동향"], "ALT Literature Survey")
    deck.section(1, "Key Insight")
    deck.bullets(deck.slide("Key Insight", "[인사이트] 요지"), [("강조 줄", True), ("보통 줄", False)])
    deck.finish(tmp_path / "out.pptx")
    slides = list(Presentation(str(tmp_path / "out.pptx")).slides)
    texts = [" ".join(sh.text_frame.text for sh in s.shapes if sh.has_text_frame) for s in slides]
    assert len(slides) == 4
    assert "해양구조물" in texts[0] and "01" in texts[1] and "Key Insight" in texts[1] and "강조 줄" in texts[2]
    assert "CREDITS" in texts[3] and "youremail" not in texts[3]
    # 양식 꾸밈 중 바깥 테두리·회색 원은 지움 (사용자 피드백 2026-10-21)
    grey = [sh for s in slides for sh in s.shapes if sh.fill.type == 1 and str(sh.fill.fore_color.rgb) == "E7E7E7"]
    assert grey == []
    fonts = {r.font.name for s in slides for sh in s.shapes if sh.has_text_frame
             for p in sh.text_frame.paragraphs for r in p.runs if r.text.strip()}
    assert fonts == {FONT}
