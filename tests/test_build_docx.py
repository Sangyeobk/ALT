import sys
from pathlib import Path

import pytest
from PIL import Image
from docx import Document
from docx.shared import Cm, Pt
from io import BytesIO

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "alt"))

from build_docx import NAVY, Report  # noqa: E402


@pytest.mark.parametrize("kind", ["table", "key_box", "flow"])
def test_table_cells_left_text_horizontally_and_center_vertically(kind):
    from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    report = Report()
    if kind == "table":
        report.table(["A", "B"], [["one", "two"]], [4, 4])
    elif kind == "key_box":
        report.key_box("Title", ["one", "two"])
    else:
        report.flow(["one", "two"])
    assert report.doc.tables[0].alignment == WD_TABLE_ALIGNMENT.LEFT
    for row in report.doc.tables[0].rows:
        for cell in row.cells:
            assert cell.vertical_alignment == WD_CELL_VERTICAL_ALIGNMENT.CENTER
            assert all(p.alignment == WD_ALIGN_PARAGRAPH.LEFT for p in cell.paragraphs)


@pytest.mark.parametrize("widths", [[4, 4], [12, 24]])
def test_table_widths_fit_body_and_match_grid_after_save(widths):
    report = Report()
    report.doc.sections[-1].left_margin = Cm(4)
    report.table(["A", "B"], [["long text " * 30, "two"]], widths)
    report.key_box("Title", ["one"])
    report.flow(["one", "two", "three"])
    stream = BytesIO()
    report.doc.save(stream)
    stream.seek(0)
    document = Document(stream)
    section = document.sections[-1]
    body_width = section.page_width - section.left_margin - section.right_margin
    for table in document.tables:
        assert table.autofit is False
        grid_widths = [column.width for column in table.columns]
        assert sum(grid_widths) <= body_width
        for row in table.rows:
            assert row.height is None
            assert [cell.width for cell in row.cells] == grid_widths
    actual = [column.width for column in document.tables[0].columns]
    assert actual[0] / actual[1] == pytest.approx(widths[0] / widths[1], abs=0.001)
    if sum(Cm(width) for width in widths) < body_width:
        assert actual == pytest.approx([Cm(width) for width in widths], abs=635)


@pytest.mark.parametrize("pixels,width_cm", [((100, 2000), 16.5), ((2000, 100), 40), ((400, 300), 5)])
def test_figure_fits_body_with_caption_and_preserves_aspect(tmp_path, pixels, width_cm):
    path = tmp_path / "figure.png"
    Image.new("RGB", pixels, "white").save(path)
    report = Report()
    section = report.doc.sections[-1]
    section.page_height = Cm(20)
    section.left_margin = Cm(4)
    report.figure(path, "Caption", width_cm)
    shape = report.doc.inline_shapes[0]
    assert shape.width <= min(Cm(width_cm), section.page_width - section.left_margin - section.right_margin)
    assert shape.height <= section.page_height - section.top_margin - section.bottom_margin - Cm(2)
    assert shape.width / shape.height == pytest.approx(pixels[0] / pixels[1], rel=0.00001)
    assert report.doc.paragraphs[-2].paragraph_format.keep_with_next
    assert report.doc.paragraphs[-1].text == "Caption"


def test_headings_use_word_heading_styles_for_outline_view():
    """장·절 제목이 Word 제목 스타일이어야 탐색 창·개요 보기에 잡힘 (사용자 피드백 2026-10-16)."""
    rpt = Report()
    rpt.heading("1. 조사 목적과 범위")
    rpt.heading("4.1 파랑하중 직접 예측", 2)
    rpt.para("본문")
    styles = [p.style.name for p in rpt.doc.paragraphs]
    assert styles == ["Heading 1", "Heading 2", "Normal"]
    h1 = rpt.doc.styles["Heading 1"]
    assert h1.font.size == Pt(15) and h1.font.color.rgb == NAVY and h1.paragraph_format.keep_with_next


def test_table_cells_break_lines_and_skip_abbreviation_expansion():
    """표 칸: '\n'마다 줄을 나누고 약어를 풀지 않음 — 풀이는 본문 첫 등장에서 (사용자 피드백 2026-10-27: 표가 복잡하고 줄바꿈 없음)."""
    rpt = Report({"CFD": "Computational Fluid Dynamics"})
    rpt.table(["항목", "내용"], [["검색어", "CFD 대리모델\nwave load"]], [3, 13])
    cell = rpt.doc.tables[0].rows[1].cells[1]
    assert [p.text for p in cell.paragraphs] == ["CFD 대리모델", "wave load"]
    rpt.para("CFD 결과")
    assert rpt.doc.paragraphs[-1].text == "CFD(Computational Fluid Dynamics) 결과"


def test_short_cite_and_no_repeated_keyword_head():
    """표·강조 상자 인용은 쪽 번호 없이 키 2개, 키워드 표시와 같은 머리말은 한 번만 (2026-10-27)."""
    from build_docx import cell_lines, cited_short, tagged
    item = {"cites": ["A2025", "B2026", "C2024"], "pages": {"A2025": 3}, "keyword": "K1"}
    assert cited_short(item) == " [A2025; B2026 외 1편]"
    kw = [{"id": "K1", "term": "계산 절감"}]
    assert tagged(item, kw, "계산 절감: 빠름") == "계산 절감: 빠름"
    assert tagged(item, kw, "빠름") == "[계산 절감] 빠름"
    assert cell_lines("검색어 3개 — 'a' / 'b'") == "검색어 3개\n'a'\n'b'"