"""발표 구성안 (사용자 결정 2026-10-21 — 양식 규칙에 맞춰 발표를 다시 씀): 검사와 양식 채우기."""

import json
import sys
from pathlib import Path

import pytest

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "alt"))
from slides_plan import check_plan, weighted_len  # noqa: E402

CATALOG = {"types": {
    "cover": {"slide": 1, "slots": {"title": {"shape": 0, "paras": ["제목"]}}},
    "items": {"slide": 2, "slots": {"body": {"shape": 1, "paras": ["a", "b", "c"], "flexible": True},
                                    "pair": {"shape": 2, "paras": ["이름", "설명"]}}},
    "figure": {"slide": 3, "slots": {"title": {"shape": 0, "paras": ["제목"]}}, "image": {"box": [0, 0, 1, 1]}},
    "table": {"slide": 4, "slots": {"title": {"shape": 0, "paras": ["제목"]}}, "table": {"box": [0, 0, 1, 1]}},
}}
LIMITS = {"cover": {"title": [{"role": "제목", "max": 10}]},
          "items": {"body": [{"role": r, "max": 20} for r in "abc"], "pair": [{"role": "이름", "max": 5}, {"role": "설명", "max": 20}]},
          "figure": {"title": [{"role": "제목", "max": 10}]}, "table": {"title": [{"role": "제목", "max": 10}]}}


def test_weighted_len_counts_latin_as_narrow():
    assert weighted_len("한글") == 2 and weighted_len("ab") == pytest.approx(1.1)


def test_check_plan_passes_good_plan(tmp_path):
    (tmp_path / "f.png").write_bytes(b"x")
    plan = {"slides": [{"type": "cover", "slots": {"title": ["파랑하중"]}},
                       {"type": "items", "slots": {"body": ["하나 [Kim2025]"], "pair": ["이름", "설명"]}},
                       {"type": "figure", "slots": {"title": ["그림"]}, "image": "f.png"},
                       {"type": "table", "slots": {"title": ["표"]}, "table": {"header": ["가", "나"], "rows": [["1", "2"]]}}]}
    assert check_plan(plan, CATALOG, LIMITS, {"Kim2025"}, tmp_path) == []


def test_check_plan_reports_each_kind_of_error(tmp_path):
    plan = {"slides": [{"type": "items", "slots": {"body": ["가" * 21], "pair": ["이름만"], "extra": ["x"]}},
                       {"type": "nope", "slots": {}},
                       {"type": "figure", "slots": {"title": ["그림 [Lee2024]"]}, "image": "없음.png"},
                       {"type": "table", "slots": {"title": ["표"]}, "table": {"header": ["가"], "rows": [["1", "2"]] * 7}}]}
    errors = " | ".join(check_plan(plan, CATALOG, LIMITS, {"Kim2025"}, tmp_path))
    for expected in ("첫 장은 cover", "한도 20자", "'pair' 문단 1개", "없는 칸 'extra'", "없는 장 종류", "그림 파일 없음",
                     "Lee2024", "표는 머리 행", "칸 수가 머리 행과 다름"):
        assert expected in errors, expected


def test_fill_template_01_from_plan(tmp_path):
    """양식 01이 있으면: 구성안의 장 종류대로 복제·채움, 장 구분 번호, 그림·표, 예시 글이 남지 않음, 글꼴 맑은 고딕."""
    from pptx import Presentation
    from build_pptx import FONT, Deck, build_talk_plan, pptx_style
    from slides_plan import load_catalog, slot_limits
    style = pptx_style("01")
    if style["path"] is None:
        pytest.skip("templates/ 양식 01 없음")
    catalog, limits = load_catalog(style), None
    limits = slot_limits(style, catalog)
    from PIL import Image
    Image.new("RGB", (400, 300), "white").save(tmp_path / "f.png")
    fill = lambda kind: {slot: [f"{l['role']}" [: max(1, int(l["max"]))] for l in ls] for slot, ls in limits[kind].items()}
    plan = {"slides": [{"type": "cover", "slots": fill("cover")}, {"type": "contents", "slots": fill("contents")},
                       {"type": "section", "slots": fill("section")}, {"type": "numbers3", "slots": fill("numbers3")},
                       {"type": "figure_side", "slots": fill("figure_side"), "image": "f.png"},
                       {"type": "section", "slots": fill("section")},
                       {"type": "table", "slots": fill("table"), "table": {"header": ["가", "나"], "rows": [["1", "2"]]}},
                       {"type": "list", "slots": {"title": ["제목"], "body": ["항목", "설명"]}}]}
    assert check_plan(plan, catalog, limits, set(), tmp_path) == []
    deck = Deck(style)
    build_talk_plan(deck, plan, catalog, tmp_path, {})
    deck.finish(tmp_path / "out.pptx")
    slides = list(Presentation(str(tmp_path / "out.pptx")).slides)
    assert len(slides) == len(plan["slides"]) + 1  # + 마무리
    texts = [" ".join(sh.text_frame.text for sh in s.shapes if sh.has_text_frame) for s in slides]
    assert "02" in texts[5] and "CREDITS" in texts[-1]
    assert any(sh.shape_type == 13 for sh in slides[4].shapes)  # 그림
    assert any(sh.has_table for sh in slides[6].shapes)
    # 바깥 테두리(채우기 없는 큰 사각형)는 지워짐 (사용자 피드백 2026-10-21) — 그림·표는 빼고 봄
    big = lambda sh: sh.width >= 0.75 * deck.prs.slide_width and sh.height >= 0.7 * deck.prs.slide_height
    frames = [sh for s in slides for sh in s.shapes
              if big(sh) and sh.shape_type not in (13, 19) and not sh.has_table and sh.fill.type != 1]
    assert frames == []
    fonts = {r.font.name for s in slides[:-1] for sh in s.shapes if sh.has_text_frame
             for p in sh.text_frame.paragraphs for r in p.runs if r.text.strip()}
    assert fonts == {FONT}


def test_check_learn_plan_needs_five_sections_and_tail_in_contents():
    from slides_plan import LEARN_TAIL, check_learn_plan
    section = {"type": "section", "slots": {"title": ["장", "설명"]}}
    good = {"slides": [{"type": "contents", "slots": {"item6": [LEARN_TAIL[0], "설명"]}}] + [section] * 5}
    assert check_learn_plan(good) == []
    bad = {"slides": [{"type": "contents", "slots": {"item6": ["용어", "설명"]}}] + [section] * 6}
    assert len(check_learn_plan(bad)) == 2


def test_fill_learn_deck_appends_tail_glossary_and_how_to_read(tmp_path):
    """배경학습 (사용자 결정 2026-10-21): 구성안 장 뒤에 장 구분 06 '용어와 읽는 법' → 용어 표(양식 표 장, 6행씩) → 읽는 법 3장."""
    from pptx import Presentation
    from build_pptx import Deck, build_learn_plan, pptx_style
    from slides_plan import LEARN_TAIL, load_catalog
    style = pptx_style("01")
    if style["path"] is None:
        pytest.skip("templates/ 양식 01 없음")
    catalog = load_catalog(style)
    plan = {"slides": [{"type": "cover", "slots": {"title": ["제목"], "subtitle": ["부제"]}}]
            + [{"type": "section", "slots": {"title": ["장", "설명"]}}] * 5}
    s = {"key_messages": [{"text": "메시지", "cites": ["Kim2025"]}],
         "background": {"glossary": [{"term": f"용어{i}", "plain": "설명"} for i in range(7)]}}
    meta = {"weights": {"relevance": 0.6, "citation": 0.2, "venue": 0.2}, "min_relevance": 4}
    deck = Deck(style)
    build_learn_plan(deck, plan, catalog, tmp_path, s, meta, {"Kim2025": {"stars": 3}})
    deck.finish(tmp_path / "learn.pptx")
    slides = list(Presentation(str(tmp_path / "learn.pptx")).slides)
    texts = [" ".join(sh.text_frame.text for sh in sl.shapes if sh.has_text_frame) for sl in slides]
    assert len(slides) == 6 + 1 + 2 + 3 + 1  # 구성안 + 장 구분 + 용어 7개 → 2장 + 읽는 법 + 마무리
    assert "06" in texts[6] and LEARN_TAIL[0] in texts[6]
    assert all(any(sh.has_table for sh in slides[i].shapes) for i in (7, 8)) and "(계속)" in texts[8]
    assert "How to Read" in texts[9]


def test_figure_side_has_body_title_and_room_for_explanation(tmp_path):
    from PIL import Image
    from pptx import Presentation
    from build_pptx import Deck, pptx_style
    from slides_plan import load_catalog, slot_limits
    style = pptx_style("01")
    if style["path"] is None:
        pytest.skip("templates/ 양식 01 없음")
    catalog = load_catalog(style)
    limits = slot_limits(style, catalog)
    Image.new("RGB", (600, 900), "white").save(tmp_path / "figure.png")
    title = "순방향 파랑력 예측"
    descriptions = ["입력 자료와 처리 순서를 설명하는 첫 번째 문단입니다.",
                    "모델을 학습하고 예측값을 구하는 두 번째 문단입니다.",
                    "그림에서 비교할 지점과 읽는 순서를 알려주는 세 번째 문단입니다."]
    points = {"type": "points3", "slots": {"title": [title], "body": ["목적", "설명", "결과", "설명", "과제", "설명"]}}
    figure = {"type": "figure_side", "slots": {"headline": [title], "text": descriptions}, "image": "figure.png"}
    deck = Deck(style)
    for spec in [points, figure]:
        deck.plan_slide(catalog["types"][spec["type"]], spec, tmp_path, 1, {"figure.png": "출처: Kim2025, Fig. 1 (p.2)"})
    deck.finish(tmp_path / "out.pptx")
    slides = list(Presentation(tmp_path / "out.pptx").slides)
    first, second = slides[1:3]
    titles = [next(sh for sh in slide.shapes if sh.has_text_frame and sh.text == title) for slide in (first, second)]
    assert [(sh.left, sh.top, sh.width, sh.height) for sh in titles][0] == [(sh.left, sh.top, sh.width, sh.height) for sh in titles][1]
    assert titles[0].text_frame.paragraphs[0].runs[0].font.size == titles[1].text_frame.paragraphs[0].runs[0].font.size
    body = next(sh for sh in second.shapes if sh.has_text_frame and descriptions[0] in sh.text)
    assert all(text in body.text for text in descriptions)
    picture = next(sh for sh in second.shapes if sh.shape_type == 13)
    assert titles[1].top + titles[1].height < min(body.top, picture.top)
    assert body.left + body.width < picture.left
    plan = {"slides": [{"type": "cover", "slots": {"title": ["조사"], "subtitle": ["범위"]}}, figure]}
    assert check_plan(plan, catalog, limits, set(), tmp_path) == []


def test_missing_catalog_file_falls_back_to_plain_build():
    from slides_plan import load_catalog
    assert load_catalog({"catalog": "templates/no_such_catalog.yaml"}) is None  # 양식만 넣고 목록은 없음 (Codex 검토 2026-10-28)
