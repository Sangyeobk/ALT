"""인용 검사: 산출물에 나온 인용 키가 모두 references.json(검증된 문헌)에 있는지 → 하나라도 없으면 실패

    python alt/check_citations.py outputs/<run>/

검사 대상: 00_summary, 01_background, 03_report, 04_presentation, 06_context (하나라도 없으면 실패) + 이 조사 맞춤 결과물(project/extras.yaml, cite_check)
"""

import argparse
import re
from pathlib import Path

from docx import Document
from pptx import Presentation

from openpyxl import load_workbook

from common import add_variant, extra_files, read_json, variant_name, write_json

KEY_PATTERN = re.compile(r"\b[A-Z][A-Za-z]+(?:19|20)\d\d[a-z]?\b")
BRACKET_PATTERN = re.compile(r"\[([^\[\]]+)\]")


def docx_text(path: Path) -> str:
    doc = Document(str(path))
    cells = [c.text for t in doc.tables for row in t.rows for c in row.cells]
    return "\n".join([p.text for p in doc.paragraphs] + cells)


def pptx_text(path: Path) -> str:
    parts = []
    for slide in Presentation(str(path)).slides:
        for shape in slide.shapes:
            if shape.has_text_frame:
                parts.append(shape.text_frame.text)
            if shape.has_table:
                parts += [c.text for row in shape.table.rows for c in row.cells]
    return "\n".join(parts)


def xlsx_text(path: Path) -> str:
    wb = load_workbook(str(path), read_only=True)
    return "\n".join(str(v) for ws in wb.worksheets for row in ws.iter_rows(values_only=True) for v in row if v is not None)


SUFFIX_READERS = {".md": lambda p: p.read_text(encoding="utf-8"), ".txt": lambda p: p.read_text(encoding="utf-8"),
                  ".docx": docx_text, ".pptx": pptx_text, ".xlsx": xlsx_text}  # 조사 맞춤 결과물 (확장자로 고름)

READERS = {"00_summary.md": lambda p: p.read_text(encoding="utf-8"), "06_context.md": lambda p: p.read_text(encoding="utf-8"),
           "03_report.docx": docx_text, "01_background.pptx": pptx_text, "04_presentation.pptx": pptx_text}


def output_texts(run_dir: Path, variant: str) -> tuple[dict[str, str], list[str]]:
    """(파일별 본문, 읽지 못한 파일). Word·PowerPoint에서 열려 잠긴 파일은 읽지 못한 것으로 보고한다."""
    texts, unreadable = {}, []
    targets = [(variant_name(base, variant), reader) for base, reader in READERS.items()]
    targets += [(name, SUFFIX_READERS[Path(name).suffix]) for name in extra_files(run_dir, variant, cite_only=True)
                if Path(name).suffix in SUFFIX_READERS]
    for name, reader in targets:
        if (run_dir / name).is_file():
            try:
                texts[name] = reader(run_dir / name)
            except Exception:  # noqa: BLE001 — 잠김·손상 모두 '확인 못 함'으로 보고
                unreadable.append(name)
    return texts, unreadable


def missing_outputs(run_dir: Path, variant: str) -> list[str]:
    """⑨에서 늘 만드는 기본 산출물 중 없는 것 — 검사 대상이 없는데 통과로 쓰지 않게 (Codex 검토 2026-10-28)."""
    return [name for name in (variant_name(base, variant) for base in READERS) if not (run_dir / name).is_file()]


def cited_keys(text: str) -> set[str]:
    """[Kim2025a; Lee2026] 형태의 괄호 안 키 + 표 칸에 단독으로 쓴 키."""
    keys = set()
    for inside in BRACKET_PATTERN.findall(text):
        keys |= set(KEY_PATTERN.findall(inside))
    for line in text.splitlines():
        if KEY_PATTERN.fullmatch(line.strip()):
            keys.add(line.strip())
    return keys


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("run_dir", type=Path)
    add_variant(parser)
    args = parser.parse_args()

    allowed = set(read_json(args.run_dir / "references.json"))
    report, all_cited = {}, set()
    texts, unreadable = output_texts(args.run_dir, args.variant)
    for name, text in texts.items():
        keys = cited_keys(text)
        all_cited |= keys
        report[name] = {"cited": len(keys), "unknown": sorted(keys - allowed)}
    unknown = sorted({k for r in report.values() for k in r["unknown"]})
    missing = missing_outputs(args.run_dir, args.variant)
    result = {"files": report, "unreadable": unreadable, "missing": missing, "unknown_keys": unknown,
              "never_cited": sorted(allowed - all_cited), "passed": not unknown and not unreadable and not missing}
    write_json(args.run_dir / variant_name("citation_check.json", args.variant), result)
    for name, r in report.items():
        print(f"[cite] {name}: 인용 키 {r['cited']}개, 목록에 없는 키 {len(r['unknown'])}개 {r['unknown'] or ''}")
    for name in unreadable:
        print(f"[cite] {name}: 읽지 못함 (Word·PowerPoint에서 열려 있으면 닫고 다시 실행) → 확인 못 함")
    if missing:
        print(f"[cite] 기본 산출물 없음 {missing} → ⑨를 먼저 끝낼 것 (확인 못 함)")
    print(f"[cite] {'통과' if result['passed'] else '실패'} — 한 번도 인용되지 않은 문헌 {len(result['never_cited'])}편")
    raise SystemExit(0 if result["passed"] else 1)


if __name__ == "__main__":
    main()
