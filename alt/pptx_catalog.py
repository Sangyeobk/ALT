"""발표 양식의 장별 칸 목록 초안 — 장 종류 목록(<양식>_catalog.yaml)을 만들 때 참고 (읽기만, 사용자 결정 2026-10-21)

    python alt/pptx_catalog.py "templates/01*.pptx" 1 3 4 ...

각 장의 도형 번호·종류·위치(cm)·글자 크기·문단 수·예시 글을 출력한다. 칸 이름·역할은 사람이(또는 Claude가) 골라 catalog에 적는다.
"""

from __future__ import annotations

import sys
from pathlib import Path

from pptx import Presentation
from pptx.util import Emu

import common  # noqa: F401 — 콘솔 한글 설정

CM = lambda v: round(Emu(v).cm, 2)


def first_size(shape) -> float | None:
    """글상자 첫 글 조각의 글자 크기 (pt) — 없으면 None (틀에서 물려받음)."""
    for p in shape.text_frame.paragraphs:
        for r in p.runs:
            if r.font.size:
                return r.font.size.pt
    return None


def describe(path: Path, numbers: list[int]) -> None:
    slides = list(Presentation(str(path)).slides)
    for n in numbers:
        slide = slides[n - 1]
        print(f"=== slide {n} ({slide.slide_layout.name})")
        for i, sh in enumerate(slide.shapes):
            kind = str(sh.shape_type).split(".")[-1].split(" ")[0]
            text = sh.text_frame.text.strip() if sh.has_text_frame else ""
            if not text and kind not in ("PICTURE", "TABLE", "CHART"):
                continue
            extra = ""
            if sh.has_text_frame and text:
                levels = [p.level for p in sh.text_frame.paragraphs if "".join(r.text for r in p.runs).strip()]
                extra = f" size={first_size(sh)} paras={len(levels)} levels={sorted(set(levels))} '{text[:50]!s}'".replace("\n", " / ")
            print(f"  [{i}] {kind:11} ({CM(sh.left)},{CM(sh.top)}) {CM(sh.width)}x{CM(sh.height)}{extra}")


if __name__ == "__main__":
    describe(next(Path.cwd().glob(sys.argv[1])), [int(n) for n in sys.argv[2:]])
