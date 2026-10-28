"""발표 구성안(slides.json): 양식 장 종류 목록 읽기, 칸별 글자 수 한도, 구성안 검사 (사용자 결정 2026-10-21 — 양식 규칙에 맞춰 다시 씀)

글자 수 한도 = 칸 폭에 들어가는 한글 글자 수 × 예시 글이 차지한 줄 수. 한글·한자 1자 = 1, 영문·숫자·기호 = 0.55 (weighted_len).
"""

from __future__ import annotations

import math
import re
import copy
from pathlib import Path

import yaml
from pptx import Presentation
from pptx.util import Emu

from check_citations import cited_keys
from common import PROJECT_DIR, load_config

PT_CM = 0.0353
LATIN_WIDTH = 0.55        # 영문·숫자 한 글자 폭 (한글 = 1)
FILL_RATIO = 0.95         # 칸을 꽉 채우지 않게 여유
DEFAULT_PT = 16.0         # 글자 크기를 틀에서 물려받아 알 수 없을 때
TABLE_MAX_ROWS, TABLE_MAX_COLS, TABLE_CELL_MAX = 6, 5, 40   # [제안] 표 장 한도 (머리 행 제외 행 수, 열 수, 칸 글자 수)
IMAGE_TYPES_KEY, TABLE_TYPES_KEY = "image", "table"


def weighted_len(text: str) -> float:
    return sum(LATIN_WIDTH if ord(ch) < 0x1100 else 1.0 for ch in text)


def load_catalog(style: dict) -> dict | None:
    path = style.get("catalog")
    if not path:
        return None
    if not (PROJECT_DIR / path).is_file():  # 양식만 넣고 장 종류 목록은 없음 → 구성안 없이 예전 방식 (Codex 검토 2026-10-28)
        print(f"[pptx] 장 종류 목록 없음 ({path}) — 구성안 없이 양식에 바로 채움")
        return None
    catalog = yaml.safe_load((PROJECT_DIR / path).read_text(encoding="utf-8"))
    layout = style.get("figure_side_layout")
    if layout and "figure_side" in catalog["types"]:
        side = catalog["types"]["figure_side"]
        title_kind = catalog["types"][layout["title_from"]]
        title_slot = title_kind["slots"]["title"]
        old_shape = side["slots"]["headline"]["shape"]
        side["slots"]["headline"] = dict(copy.deepcopy(title_slot), shape=old_shape,
                                            source=[title_kind["slide"], title_slot["shape"]], is_title=True)
        side["slots"]["text"].update(box=layout["text_box"], paras=layout["text_roles"], repeat=True,
                                      flexible=True, lines=layout["text_lines"],
                                      space_after_pt=layout["text_space_after_pt"])
        side["image"]["box"] = layout["image_box"]
        side["use"] = "상단 본문 제목 + 왼쪽 그림 소개·읽는 순서·확인할 지점 + 오른쪽 원문 그림"
    return catalog


def _para_size(paragraph) -> float:
    for r in paragraph.runs:
        if r.font.size:
            return r.font.size.pt
    return DEFAULT_PT


def slot_limits(style: dict, catalog: dict) -> dict:
    """장 종류 → 칸 → [{role, max}] (글이 있는 문단마다). 양식 파일의 예시 글로 계산."""
    slides = list(Presentation(str(style["path"])).slides)
    out = {}
    for name, kind in catalog["types"].items():
        shapes = list(slides[kind["slide"] - 1].shapes)
        out[name] = {}
        for slot, spec in (kind.get("slots") or {}).items():
            shape = slides[spec["source"][0] - 1].shapes[spec["source"][1]] if spec.get("source") else shapes[spec["shape"]]
            width = spec["box"][2] if spec.get("box") else (spec.get("widen") or [Emu(shape.width).cm])[0]
            filled = [p for p in shape.text_frame.paragraphs if "".join(r.text for r in p.runs).strip()]
            if spec.get("repeat"):
                filled = [filled[0]] * len(spec["paras"])
            limits = []
            for role, p in zip(spec["paras"], filled):
                size = _para_size(p)
                sample = "".join(r.text for r in p.runs)
                per_line = Emu(shape.width).cm / (size * PT_CM)          # 예시 글이 원래 칸에서 몇 줄이었나
                lines = spec.get("lines") or max(1, math.ceil(weighted_len(sample) / per_line - 0.15))  # 목록에서 줄 수를 정할 수 있음
                limits.append({"role": role, "max": int(width / (size * PT_CM) * lines * FILL_RATIO)})
            out[name][slot] = limits
    return out


def check_plan(plan: dict, catalog: dict, limits: dict, ref_keys: set[str], run_dir: Path) -> list[str]:
    """구성안 검사 — 오류 문장 목록 (비어 있으면 통과)."""
    errors = []
    slides = plan.get("slides") or []
    if not slides or slides[0].get("type") != "cover":
        errors.append("첫 장은 cover여야 함")
    for n, spec in enumerate(slides, 1):
        kind_name = spec.get("type")
        kind = catalog["types"].get(kind_name)
        where = f"{n}장({kind_name})"
        if kind is None:
            errors.append(f"{where}: 없는 장 종류")
            continue
        given = spec.get("slots") or {}
        for slot, spec_slot in (kind.get("slots") or {}).items():
            texts = given.get(slot)
            if isinstance(texts, str):
                texts = [texts]
            lim = limits[kind_name][slot]
            if not texts:
                errors.append(f"{where}: 칸 '{slot}' 비어 있음")
                continue
            if len(texts) > len(lim) or (len(texts) < len(lim) and not spec_slot.get("flexible")):
                errors.append(f"{where}: 칸 '{slot}' 문단 {len(texts)}개 — {len(lim)}개여야 함" + (" 이하" if spec_slot.get("flexible") else ""))
            for text, l in zip(texts, lim):
                if weighted_len(text) > l["max"]:
                    errors.append(f"{where}: '{slot}'의 {l['role']} {weighted_len(text):.0f}자 > 한도 {l['max']}자: {text[:30]}…")
        for slot in given:
            if slot not in (kind.get("slots") or {}):
                errors.append(f"{where}: 없는 칸 '{slot}'")
        if IMAGE_TYPES_KEY in kind and not (run_dir / str(spec.get("image", ""))).is_file():
            errors.append(f"{where}: 그림 파일 없음 '{spec.get('image')}'")
        if TABLE_TYPES_KEY in kind:
            table = spec.get("table") or {}
            header, rows = table.get("header") or [], table.get("rows") or []
            if not header or not rows or len(rows) > TABLE_MAX_ROWS or len(header) > TABLE_MAX_COLS:
                errors.append(f"{where}: 표는 머리 행 + 1–{TABLE_MAX_ROWS}행, 최대 {TABLE_MAX_COLS}열")
            for row in [header] + rows:
                if len(row) != len(header):
                    errors.append(f"{where}: 표 행의 칸 수가 머리 행과 다름")
                for cell in row:
                    if weighted_len(str(cell)) > TABLE_CELL_MAX:
                        errors.append(f"{where}: 표 칸 {weighted_len(str(cell)):.0f}자 > {TABLE_CELL_MAX}자: {str(cell)[:30]}…")
        text = " ".join(str(t) for v in given.values() for t in (v if isinstance(v, list) else [v]))
        text += " " + " ".join(str(c) for row in ((spec.get("table") or {}).get("rows") or []) for c in row)
        unknown = cited_keys(text) - ref_keys
        if unknown:
            errors.append(f"{where}: 참고문헌에 없는 인용 키 {sorted(unknown)}")
    return errors


LEARN_FIGURE_PREFIX = "figures/slide_"   # 배경학습에는 발표용 단순화 그림만


def figure_choices(run_dir: Path, synthesis: dict) -> list[dict]:
    """구성안에 쓸 수 있는 그림 (있는 파일만). 원문 그림은 출처를 코드가 그림 아래에 붙임."""
    # 보고서용 그림(흐름도·계통도·한계→대응)은 붙이지 않고 발표 스타일로 단순화한 그림(make_figures slide_*)을 씀
    # (사용자 피드백 2026-10-21) — 셋 다 가로로 넓음(2.2:1) → figure_full
    made = [("02_literature_map.png", "문헌지도 — 후보 문헌 분포 (가로축 × 접근법, 원 크기 = 총점)"),
            ("figures/prisma.png", "문헌 선별 흐름도 (PRISMA 형식) — 세로로 긴 그림"),
            ("figures/distribution.png", "논문 수 분포 — 대륙·가로축·접근법"),
            ("figures/slide_flow.png", "기술 흐름 단계 — 단계마다 기존 방법·핵심논문 편수·분야 수 (넓음)"),
            ("figures/slide_limits.png", "기존 방법의 한계 → 역할 → 핵심논문 편수 (넓음)"),
            ("figures/slide_tree.png", "기술 계통도 — 주제 → 역할 → 분야·편수 (넓음)")]
    out = [{"file": f, "what": what} for f, what in made if (run_dir / f).is_file()]
    for fig in synthesis.get("figures", []):
        if (run_dir / fig["file"]).is_file():
            out.append({"file": fig["file"], "what": f"원문 그림 {fig['key']} {fig['fig']} (p.{fig['page']}) — {fig.get('caption', '')}",
                        "subtopic": fig.get("subtopic")})
    return out


def figure_sources(synthesis: dict) -> dict:
    """원문 그림 파일 → 그림 아래 출처 줄."""
    return {fig["file"]: f"출처: {fig['key']}, {fig['fig']} (p.{fig['page']})" for fig in synthesis.get("figures", [])}


SLIDES_FILE = "slides.json"
INPUT_FILE = "slides_input.json"
# 발표 종류 → (구성안 입력, 구성안) — 01_background도 같은 방식으로 다시 씀 (사용자 결정 2026-10-21)
DECK_FILES = {"talk": (INPUT_FILE, SLIDES_FILE), "learn": ("slides_learn_input.json", "slides_learn.json")}
# 배경학습 마지막 장 구분 — 용어 표·읽는 법은 코드가 붙임 (사용자 결정 2026-10-21: 용어는 양식 표 장에, 읽는 법은 지금 그림 그대로)
LEARN_TAIL = ["용어와 읽는 법", "용어 풀이와 결과물 읽는 법"]
LEARN_SECTIONS = 5   # 구성안이 쓰는 장 구분 수 — 목차 6번째는 LEARN_TAIL


def check_learn_plan(plan: dict) -> list[str]:
    """배경학습 구성안 추가 검사: 장 구분 5개, 목차 6번째 = 코드가 붙이는 마지막 장."""
    errors = []
    slides = plan.get("slides") or []
    n = sum(spec.get("type") == "section" for spec in slides)
    if n != LEARN_SECTIONS:
        errors.append(f"장 구분 {n}개 — 배경학습은 {LEARN_SECTIONS}개 (6번째 '{LEARN_TAIL[0]}'는 코드가 붙임)")
    contents = next((spec for spec in slides if spec.get("type") == "contents"), None)
    if contents is None or (contents.get("slots", {}).get("item6") or [""])[0] != LEARN_TAIL[0]:
        errors.append(f"목차 item6의 장 이름은 '{LEARN_TAIL[0]}'여야 함")
    return errors
