"""문헌조사 보고서 → 03_보고서.docx (모든 인용은 참고문헌 키와 연결)

    python alt/build_docx.py outputs/<run>/
장 구성·순서: alt/report_guide.md, config output.report_sections (개인 설정이 우선).
양식: 맑은 고딕, 제목 남색(002060), 개조식, 핵심 빨강 굵게, 표는 회색 머리줄.
"""

import argparse
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor, Twips

from common import (PRISMA_NOTE, Abbreviations, verbatim_strings, add_variant, cite, display_names, keyword_chunks, keyword_tag, tech_keywords, load_config, load_topic, map_axis, read_json, review_info,
                    report_title, role_label, scope_history, sections, topic_line, variant_name)
from export_references import EXCEL_SHEETS
from report_facts import (ADJACENT_NOTE, adjacent_rows, claim_basis, apa_reference, apa_sort_key, excluded_comparison, excluded_rows, fulltext_label, method_lines,
                          role_of)

FONT = "맑은 고딕"
NAVY = RGBColor(0x00, 0x20, 0x60)
RED = RGBColor(0xFF, 0x00, 0x00)
HEAD_MAX_CHARS = load_config()["output"]["head_max_chars"]   # '머리말: 설명'의 머리말로 볼 최대 글자 수
CITE_SHOWN = load_config()["output"]["cite_shown"]           # 본문에 보이는 인용 키 수 (나머지는 '외 N편', 전체는 엑셀)
TABLE_CITE_SHOWN = load_config()["output"]["table_cite_shown"]  # 표·강조 상자 안 인용 키 수 (쪽 번호 없이)
TABLE_KEEP_ROWS = load_config()["output"]["table_keep_rows"]   # 이 행 수 이하 표는 쪽에 통째로 (넘으면 머리 행 반복)
TECH_COLOR = RGBColor.from_string(load_config()["output"]["tech_keyword_color"])  # 기술 핵심어 글자색
FIGURE_WIDTH_CM = 16.5
FIGURE_CAPTION_SPACE_CM = 2.0  # 캡션과 그림 문단 간격을 위한 배치 여유 (실제 줄 수 측정값 아님)
EMU_PER_TWIP = 635  # Word 표 폭 저장 단위: 1/20 pt


HEADING_LOOK = {1: (15, 14), 2: (12, 8)}  # 제목 수준 → (글자 pt, 앞 간격 pt)


class Report:
    def __init__(self, abbreviations: dict | None = None, protected: list[str] | None = None):
        self.abbr = Abbreviations(abbreviations, protected)  # 약어는 문서에서 처음 나올 때 풀어 씀 (입력 원문·검색어는 그대로)
        self.tech: list[str] = []  # 기술 핵심어 (build에서 요약의 tech_keywords로 채움)
        self.doc = Document()
        for section in self.doc.sections:
            section.left_margin = section.right_margin = Cm(2.2)
            section.top_margin = section.bottom_margin = Cm(2.0)
        normal = self.doc.styles["Normal"]
        normal.font.name, normal.font.size = FONT, Pt(10.5)
        normal.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
        normal.paragraph_format.space_after = Pt(3)
        normal.paragraph_format.line_spacing = 1.25
        # 장·절 제목은 Word 제목 스타일 → 탐색 창·개요 보기에 잡힘 (사용자 피드백 2026-10-16). 모양은 KR 양식
        for level, (size, before) in HEADING_LOOK.items():
            style = self.doc.styles[f"Heading {level}"]
            rfonts = style.element.get_or_add_rPr().get_or_add_rFonts()
            for attr in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
                rfonts.attrib.pop(qn(attr), None)  # 테마 글꼴이 지정 글꼴보다 앞서므로 지움
            style.font.name, style.font.size, style.font.bold, style.font.color.rgb = FONT, Pt(size), True, NAVY
            rfonts.set(qn("w:eastAsia"), FONT)
            style.paragraph_format.space_before, style.paragraph_format.space_after = Pt(before), Pt(3)
            style.paragraph_format.keep_with_next = True

    def run(self, p, text, size=None, bold=False, color=None, mark=True, expand=True):
        """글 한 토막. 기술 핵심어(self.tech)는 초록색 글자로 (mark=False인 제목·표지는 제외). 마지막 조각을 돌려줌.
        expand=False면 약어를 풀지 않음 — 표 칸 (사용자 결정 2026-10-27: 풀이는 본문 첫 등장·부록 약어표에서)."""
        r = None
        for piece, is_tech in keyword_chunks(self.abbr(text) if expand else text, self.tech if mark else []) or [("", False)]:
            r = p.add_run(piece)
            r.font.name = FONT
            r._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
            if size:
                r.font.size = Pt(size)
            r.bold = bold
            if is_tech:
                r.font.color.rgb = TECH_COLOR
            elif color:
                r.font.color.rgb = color
        return r

    def para(self, text="", size=None, bold=False, color=None, align=None, mark=True):
        p = self.doc.add_paragraph()
        if align:
            p.alignment = align
        self.run(p, text, size, bold, color, mark)
        return p

    def heading(self, text, level=1):
        p = self.doc.add_paragraph(style=f"Heading {level}")
        self.run(p, text, bold=True, mark=False)  # 크기·색·간격은 스타일이 정함 (run()은 굵기를 직접 지정하므로 굵게 명시)

    def bullet(self, text, emphasis=False, level=0):
        """개조식 한 항목. '머리말: 설명' 형식이면 머리말을 굵게 (읽기 쉽게)."""
        p = self.doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(0.5 + 0.6 * level)
        p.paragraph_format.first_line_indent = Cm(-0.4)
        p.paragraph_format.space_after = Pt(2)
        self.run(p, "• " if level == 0 else "– ")
        head, sep, body = text.partition(": ")
        if sep and len(head) <= HEAD_MAX_CHARS:
            self.run(p, head + ": ", bold=True, color=RED if emphasis else None)
            self.run(p, body, color=RED if emphasis else None)
        else:
            self.run(p, text, bold=emphasis, color=RED if emphasis else None)

    def lead(self, text):
        """장·절 첫 줄 요지 한 문장 (굵게)."""
        if text:
            p = self.para(text, 11, True)
            p.paragraph_format.space_after = Pt(6)

    def shade(self, cell, fill):
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:color"), "auto")
        shd.set(qn("w:fill"), fill)
        cell._element.get_or_add_tcPr().append(shd)

    def keep_together(self, t, whole: bool) -> None:
        """표가 쪽 끝에서 잘리지 않게 (사용자 피드백 2026-10-19: 핵심 키워드 상자가 두 쪽으로 나뉨). 파이썬은 쪽 위치를 모르므로
        Word 설정으로: 행은 나누지 않음, whole이면 마지막 행 앞까지 '다음 단락과 함께' → 모자라면 표째 다음 쪽으로.
        긴 표는 쪽을 넘기되 머리 행을 다음 쪽에서 반복."""
        for i, row in enumerate(t.rows):
            tr_pr = row._tr.get_or_add_trPr()
            tr_pr.append(OxmlElement("w:cantSplit"))
            if i == 0 and not whole:
                tr_pr.append(OxmlElement("w:tblHeader"))
            if whole and i < len(t.rows) - 1:
                for cell in row.cells:
                    for p in cell.paragraphs:
                        p.paragraph_format.keep_with_next = True
        if whole:  # 마지막 행 안에서도 줄끼리 붙여 둠 (칸 하나짜리 강조 상자)
            for cell in t.rows[-1].cells:
                for p in cell.paragraphs[:-1]:
                    p.paragraph_format.keep_with_next = True

    def format_table(self, table, widths=None):
        """모든 표의 정렬·폭을 통일하고 본문 폭 안에서 열 비율을 유지한다."""
        section = self.doc.sections[-1]
        body_width = section.page_width - section.left_margin - section.right_margin
        requested = [Cm(width) for width in widths] if widths is not None else [body_width / len(table.columns)] * len(table.columns)
        scale = min(1, body_width / sum(requested))
        # twip 단위로 내림하여 저장 시 반올림 때문에 총폭이 본문보다 커지지 않게 한다.
        fitted = [Twips(int(width * scale / EMU_PER_TWIP)) for width in requested]
        table.alignment = WD_TABLE_ALIGNMENT.LEFT
        table.autofit = False
        for column, width in zip(table.columns, fitted):
            column.width = width
        for row in table.rows:
            for cell, width in zip(row.cells, fitted):
                cell.width = width
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                for paragraph in cell.paragraphs:
                    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT

    def table(self, header, rows, widths, size=9):
        t = self.doc.add_table(rows=1, cols=len(header))
        t.style = "Table Grid"
        for i, text in enumerate(header):
            cell = t.rows[0].cells[i]
            cell.text = ""
            self.run(cell.paragraphs[0], text, size, True)
            self.shade(cell, "D9D9D9")
        for row in rows:
            cells = t.add_row().cells
            for i, text in enumerate(row):
                cells[i].text = ""
                # 칸 안 줄바꿈: '\n'마다 문단 하나, 약어는 풀지 않음 (사용자 피드백 2026-10-27: 표가 복잡하고 줄바꿈 없음)
                for k, line in enumerate(str(text if text is not None else "").split("\n")):
                    p = cells[i].paragraphs[0] if k == 0 else cells[i].add_paragraph()
                    p.paragraph_format.space_after = Pt(0)
                    self.run(p, line, size, expand=False)
        self.format_table(t, widths)
        self.keep_together(t, whole=len(rows) <= TABLE_KEEP_ROWS)
        self.doc.add_paragraph()

    def key_box(self, title, lines, numbered=True):
        t = self.doc.add_table(rows=1, cols=1)
        t.style = "Table Grid"
        cell = t.rows[0].cells[0]
        self.shade(cell, "FFF2CC")
        cell.text = ""
        self.run(cell.paragraphs[0], title, 11, True, NAVY)
        for i, line in enumerate(lines, 1):
            self.run(cell.add_paragraph(), f"{i}. {line}" if numbered else line, 10.5, True)
        self.format_table(t)
        self.keep_together(t, whole=True)
        self.doc.add_paragraph()

    def labeled(self, text, emphasis=False):
        """부록용 짧은 글 한 문단: '머리말: 설명'이면 머리말만 굵게 (개조식 기호 없음)."""
        p = self.doc.add_paragraph()
        head, sep, body = text.partition(": ")
        if sep:
            self.run(p, head + ": ", bold=True, color=RED if emphasis else None)
        self.run(p, body if sep else text)

    def figure(self, path: Path, caption: str, width_cm=FIGURE_WIDTH_CM):
        if path.is_file():
            section = self.doc.sections[-1]
            body_width = section.page_width - section.left_margin - section.right_margin
            max_height = section.page_height - section.top_margin - section.bottom_margin - Cm(FIGURE_CAPTION_SPACE_CM)
            picture = self.doc.add_picture(str(path), width=Cm(width_cm))
            scale = min(1, body_width / picture.width, max_height / picture.height)
            picture.width, picture.height = int(picture.width * scale), int(picture.height * scale)
            self.doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
            self.doc.paragraphs[-1].paragraph_format.keep_with_next = True  # 캡션과 같은 쪽
        self.para(caption, 9, False, None, WD_ALIGN_PARAGRAPH.CENTER)

    def flow(self, steps: list[str]):
        """공학적 흐름: 흰 상자 → 화살표, 마지막 단계만 노랑 (KR 흐름도 양식)."""
        cells = [c for step in steps for c in (step, "→")][:-1]
        t = self.doc.add_table(rows=1, cols=len(cells))
        for i, text in enumerate(cells):
            cell = t.rows[0].cells[i]
            cell.text = ""
            self.run(cell.paragraphs[0], text, 10, True, NAVY if text == "→" else None)
            if text != "→":
                self.shade(cell, "FFF2CC" if i == len(cells) - 1 else "F2F2F2")
        self.format_table(t)
        self.doc.add_paragraph()


# ---------- 장 (id는 report_guide.md · config output.report_sections) ----------
def cited(item) -> str:
    return cite(item, CITE_SHOWN)


def cited_short(item) -> str:
    """표·강조 상자용 인용: 쪽 번호 없이 대표 키만 (사용자 결정 2026-10-27, config output.table_cite_shown)."""
    return cite(item, TABLE_CITE_SHOWN, with_pages=False)


def tagged(item: dict, keywords: list[dict], text: str) -> str:
    """핵심 키워드 표시 + 글 — 글이 이미 그 키워드 머리말로 시작하면 표시를 빼서 같은 말이 두 번 나오지 않게 (2026-10-27)."""
    tag = keyword_tag(item, keywords)
    return text if tag and text.startswith(tag[1:-2]) else tag + text


def cell_lines(text: str) -> str:
    """한 줄로 이어 붙인 목록·덧붙임을 칸 안 줄로 나눔 (' / ' 목록, ' — ' 덧붙임)."""
    return text.replace(" / ", "\n").replace(" — ", "\n")


def split_head(text: str) -> list[str]:
    """'머리말: 설명' → [머리말, 설명] (부록 표의 두 칸). 머리말이 없으면 ['', 글]."""
    head, sep, body = text.partition(": ")
    return [head, body] if sep else ["", text]


def sec_purpose(rpt, c):
    sc = c.s["scope"]
    # 확정 주제 + 범위가 정해진 과정을 '구분 | 내용' 표로 (사용자 피드백 2026-10-16·19 — 발표 목적 장과 같은 모양)
    asked = topic_line(c.run_dir, c.refs).partition(": ")[2]
    history = scope_history(c.run_dir)
    rows = [*[[label, text] for label, text in history if label == "입력 원문"], *([["조사 주제", asked]] if asked else []),
            *[[label, text] for label, text in history if label != "입력 원문"]]
    if rows:
        rpt.table(["구분", "내용"], rows, [3.0, 13.5], size=9)
    rpt.bullet("연구 문제: " + sc["question"], emphasis=True)
    if c.s.get("keywords"):  # 핵심 키워드: 처음 소개해 결론까지 이어 감 (report_guide '핵심 키워드와 용어 통일')
        rpt.key_box("핵심 키워드 : 보고서 전체에 쓰는 판단 기준",
                    [f"{k['term']}: {k['plain']} — {k['question']}" for k in c.s["keywords"]])
    if "landscape" not in sections("report", c.run_dir):  # 흐름도는 연구 접근 장에 한 번만
        rpt.flow(sc["flow"])
    if sc.get("problem_types"):
        rpt.table(["문제 유형", "입력", "출력", "왜 필요한가"],
                  [[t["name"], t["input"], t["output"], t["why"]] for t in sc["problem_types"]], [3.1, 4.2, 4.2, 5.0], size=9)
    if sc.get("framing"):
        rpt.bullet("비교 원칙: " + sc["framing"])
    rpt.bullet("핵심 범위: " + sc["core"])
    rpt.bullet("활용·확장: " + sc["extension"])
    rpt.bullet("제외: " + sc["excluded"])


def sec_summary(rpt, c):
    """핵심 인사이트: 무엇을 조사했나 한 줄(목적 장이 앞에 없을 때만) → 주요 발견 → 해석할 때 유의점."""
    if "purpose" not in sections("report", c.run_dir):  # 목적 장이 맨 앞이면 연구 문제가 이미 나옴 — 반복 금지
        rpt.bullet("조사 질문: " + c.s["scope"]["question"])
    rpt.key_box("Key Message : 주요 발견", [tagged(m, c.kw, m["text"]) + cited_short(m) for m in c.s["key_messages"]])
    rpt.heading("해석할 때 유의점", 2)
    for item in c.s["limitations"]:
        rpt.bullet(item["text"] + cited(item))


def sec_method(rpt, c):
    """부록: 핵심(강조 항목)만 짧은 글, 나머지 사실은 '항목 | 내용' 표 (사용자 결정 2026-10-15)."""
    lines = method_lines(c.meta, c.records, c.config, c.level_of, load_topic(c.run_dir))
    for text, emphasis in lines:
        if emphasis:
            rpt.labeled(text)
    rpt.table(["항목", "내용"], [[head, cell_lines(body)] for head, body in (split_head(text) for text, emphasis in lines if not emphasis)],
              [3.2, 13.3], size=9)
    for text in PRISMA_NOTE:
        rpt.labeled(text)
    rpt.figure(c.run_dir / "figures" / "prisma.png", f"그림 {c.fig()}. 문헌 선별 흐름도 (PRISMA 흐름도 형식 참조)", 12)


def sec_landscape(rpt, c):
    """연구 접근과 문헌 분포: 흐름도+문헌(배경 이론 위 연구 위치) → 한계→대응(세 접근) → 분야 → 지도.
    그림이 없으면(예전 요약 형식) 흐름 상자·비교표로 대신한다."""
    flow_png, roles_png = c.run_dir / "figures" / c.vname("theory_flow.png"), c.run_dir / "figures" / c.vname("limit_response.png")
    ap = display_names(c.run_dir, c.s, c.meta)["approach"]
    if flow_png.is_file():
        rpt.figure(flow_png, f"그림 {c.fig()}. 기술 흐름에서 {ap} 연구의 위치 —기존 방법은 [일반 개념], "
                   "편수는 핵심논문, 학습·검증 자료는 초록·본문 판독(중복 포함)")
    else:
        rpt.flow(c.s["scope"]["flow"])
    if roles_png.is_file():
        rpt.figure(roles_png, f"그림 {c.fig()}. 기존 방법의 한계와 {ap}의 대응")
        if c.s.get("approach_roles_note"):
            rpt.bullet("읽는 법: " + c.s["approach_roles_note"])
    elif c.s.get("approach_roles"):
        rpt.table(["접근(역할)", "예측 대상", "주로 쓰는 자료", "확인해야 할 성능", "해당 분야"],
                  [[a["role"], a["target"], a["data"], a["check"], a.get("subtopics", "")] for a in c.s["approach_roles"]],
                  [2.8, 3.6, 3.4, 3.8, 2.9], size=9)
        if c.s.get("approach_roles_note"):
            rpt.bullet("읽는 법: " + c.s["approach_roles_note"])
    tree_png = c.run_dir / "figures" / c.vname("tech_tree.png")
    if tree_png.is_file():  # 기술 계통도가 분야 표를 대신함 (사용자 결정 2026-10-16 — 같은 내용 반복 줄임)
        rpt.figure(tree_png, f"그림 {c.fig()}. 기술 계통도 — 역할 → 분야(핵심논문 편수·대표 연구) → 대표 기법")
    else:
        rpt.table(["분야", "역할", "편수", "다루는 문제"],
                  [[t["name"], role_label(t, c.s), len(t["papers"]), t["problem"]] for t in c.s["subtopics"]],
                  [4.2, 1.8, 1.2, 9.3], size=9)
    floor = c.config["output"]["map_min_relevance"]
    pool = sum(r["relevance"] >= floor and not r["exclude"] for r in c.records)
    rpt.figure(c.run_dir / "02_문헌지도.png", f"그림 {c.fig()}. 문헌지도 — 관련성 {floor}점 이상 후보 {pool}편 (남색 핵심논문, ◆ 문헌조사논문, 회색 기타). "
               "원 크기는 조사 우선순위 점수이며 품질 판정이 아님")
    dist_png = c.run_dir / "figures" / "distribution.png"
    if dist_png.is_file():  # 논문 수 분포 — 대륙·축·접근법 (사용자 피드백 2026-10-19)
        rpt.figure(dist_png, f"그림 {c.fig()}. 논문 수 분포 — 문헌지도와 같은 후보, 남색은 핵심논문. 대륙은 교신저자 소속(없으면 1저자 소속)")
    for item in c.s["map_observations"]:
        rpt.bullet(item["text"] + cited(item))


def sec_subtopics(rpt, c):
    for i, t in enumerate(c.s["subtopics"], 1):
        rpt.heading(f"{c.number}.{i} {t['name']} [{role_label(t, c.s)} · {len(t['papers'])}편]", 2)
        # 공학 보고서 전개 (사용자 결정 2026-10-15): 요지 한 줄 → (1) 목적과 방법 → (2) 주요 결과 → (3) 종합 및 남은 과제
        rpt.lead(t.get("lead") or t.get("synthesis", ""))
        rpt.para("(1) 연구 목적과 방법", 10.5, True, NAVY)
        rpt.bullet("목적: " + t["problem"])
        if t.get("approaches"):
            rpt.bullet("방법: " + t["approaches"])
        rpt.bullet("입력 → 출력: " + t["io"])
        if t.get("link"):
            rpt.bullet("핵심 범위와의 관계: " + t["link"])
        rpt.para("(2) 주요 결과", 10.5, True, NAVY)
        for item in t["reported"][:2]:
            rpt.bullet(item["text"] + cited(item))
        rpt.para("(3) 종합 및 남은 과제", 10.5, True, NAVY)
        if t.get("lead") and t.get("synthesis"):
            rpt.bullet("종합: " + t["synthesis"])
        if t["unverified"]:
            rpt.bullet("남은 과제: " + " / ".join(t["unverified"][:2]), emphasis=True)
        for fig in [f for f in c.s.get("figures", []) if f["subtopic"] == t["id"]]:  # 전문 모드: 원문 그림 (출처 표시)
            rpt.figure(c.run_dir / fig["file"], f"그림 {c.fig()}. {fig['caption']} — 원문 출처: {fig['key']}, {fig['fig']} (p.{fig['page']})", 12)
    extra = len(c.s["subtopics"])
    if c.s.get("cross_synthesis"):
        extra += 1
        rpt.heading(f"{c.number}.{extra} 분야 간 종합 — 같은 기법이 분야마다 하는 역할", 2)
        for item in c.s["cross_synthesis"]:
            rpt.bullet(keyword_tag(item, c.kw) + item["text"] + cited(item))
    if c.s.get("keyword_matrix") and c.kw:  # 분야 × 핵심 키워드: 결론의 키워드 판단이 여기서 나옴
        extra += 1
        rpt.heading(f"{c.number}.{extra} 분야 × 핵심 키워드", 2)
        rpt.lead((c.s.get("leads") or {}).get("matrix", ""))
        width = 12.3 / len(c.kw)
        rpt.table(["분야"] + [k["term"] for k in c.kw],
                  [[t["name"]] + [c.s["keyword_matrix"].get(t["id"], {}).get(k["id"], "미확인") for k in c.kw] for t in c.s["subtopics"]],
                  [4.2] + [width] * len(c.kw), size=9)
    if c.s.get("metrics_table"):  # 전문 모드: 본문 수치 비교
        rpt.heading(f"{c.number}.{extra + 1} 수치 비교표 (각 연구의 검증 조건에서 보고된 값)", 2)
        rpt.table(["인용 키", "지표", "값", "조건", "쪽"],
                  [[m["key"], m["metric"], m["value"], m["condition"], f"p.{m['page']}"] for m in c.s["metrics_table"]],
                  [2.6, 2.4, 2.2, 7.4, 1.4], size=8)


def sec_issues(rpt, c):
    rpt.para("순위는 연구 문제에 비춘 이 조사의 판단이며, 분야의 합의가 아님.", 9)
    for i in sorted(c.s["issues"], key=lambda i: i.get("priority", 99)):
        rpt.bullet(f"{i.get('priority', '-')}순위: {keyword_tag(i, c.kw)}{i['observation']}{cited(i)}")
        rpt.bullet("해석상 한계: " + i["limitation"], level=1)
        rpt.bullet("열린 질문: " + i["question"], level=1)
    rpt.figure(c.run_dir / "figures" / "gap_matrix.png", f"그림 {c.fig()}. 후보 문헌 분포 (접근법 × {c.axis['name']}, 중복 포함) — "
               "· 는 조사 표본에서 확인되지 않은 조합이며 연구가 없다는 뜻이 아님", 15)
    rpt.heading("연구 방향 — 관찰된 접근과 가설", 2)
    rpt.lead((c.s.get("leads") or {}).get("outlook", ""))
    for item in sorted(c.s["outlook"], key=lambda o: o.get("kind") != "관찰"):
        rpt.bullet(f"[{item.get('kind', '')}] {keyword_tag(item, c.kw)}{item['text']}{cited(item)}")


def sec_conclusion(rpt, c):
    """결론: 핵심 키워드마다 한 줄 판단 + 종합 (앞 장의 분야×키워드 표·쟁점을 요약). 예전 형식(kind)도 그대로 보임."""
    term = {k["id"]: k["term"] for k in c.kw}
    for p in c.s["conclusion"].get("points", []):
        head = term.get(p.get("keyword")) or p.get("kind", "")
        rpt.bullet(f"{head}: {p['text']}{cited(p)}" if head else p["text"] + cited(p), emphasis=p.get("kind") == "한계")
    if c.s["conclusion"].get("closing"):
        rpt.key_box("종합", [c.s["conclusion"]["closing"]], numbered=False)


def sec_appendix_next(rpt, c):
    """조사 과정 기록: 에이전트가 다음 단계(전문 분석)에서 볼 대상 — 독자용 읽기 목록이 아님. 한 문단 + 표 2개."""
    priority = c.s["conclusion"]["fulltext_priority"]
    done = sum(c.level_of.get(p["key"]) == "전문" for p in priority)
    rpt.para(f"전문 분석 대상 {len(priority)}편 중 {done}편은 본문을 읽어 위 판단에 반영함" if done else
             "초록 기반 조사이며, 아래 대상의 본문은 아직 분석하지 않음", 10)
    if c.s.get("next_steps"):
        rows = [split_head(i["text"]) for i in c.s["next_steps"]]
        rpt.table(["다음 단계", "내용"], [[head, body + cited_short(i)] for (head, body), i in zip(rows, c.s["next_steps"])],
                  [3.2, 13.3], size=9)
    rpt.table(["인용 키", "선정 이유", "본문에서 볼 질문", "상태"],
              [[p["key"], p["reason"], "\n".join(p["questions"]),
                "본문 반영" if c.level_of.get(p["key"]) == "전문" else f"전문 {fulltext_label(c.refs[p['key']])}"] for p in priority],
              [2.6, 5.8, 5.8, 2.3], size=8.5)


def sec_appendix_excel(rpt, c):  # 안내 문장은 장 요지(leads)가 맡음 — 같은 말 반복 금지
    rpt.table(["시트", "내용"], [[name, desc] for name, desc in EXCEL_SHEETS], [3.6, 12.9], size=9)


def sec_appendix_methods(rpt, c):
    role = role_of(c.s)
    rpt.table(["인용 키", "역할·관련성", "검토", "해결 문제", "방법", "검증 근거", "구조물"],
              [[m["key"], f"{role.get(m['key'], '')} · {c.refs[m['key']]['relevance']}점", c.level_of.get(m["key"], ""),
                m.get("problem", ""), m["method"], m["validation"], m["structure"]] for m in c.s["method_comparison"]],
              [2.2, 1.6, 1.1, 3.2, 3.4, 3.2, 1.8], size=8)


def sec_appendix_excluded(rpt, c):
    rows = excluded_rows(c.records, c.axis)
    rpt.bullet(f"관련성 기준은 통과했지만 총점 미달·상한 초과로 빠진 {len(rows)}편 — 선정 편수 제한이 결론을 바꾸는지 확인하기 위한 표", emphasis=True)
    rpt.table(["항목", "핵심논문", "선정 밖"], excluded_comparison(c.records, c.axis), [6.0, 2.5, 2.5], size=9)
    if rows:
        rpt.table(["순위", "제목", "연도", "종류", "인용", "관련성", "총점", c.axis["name"], "이유"], rows,
                  [1.1, 5.4, 1.2, 1.5, 1.1, 1.3, 1.3, 2.3, 1.5], size=8)


def sec_appendix_counts(rpt, c):
    claims = c.s.get("count_claims", [])
    if not claims:
        rpt.bullet("편수 주장 근거가 기록되지 않음 (요약 형식 이전 버전)")
        return
    rpt.para("자동 집계 = 점수·분류 기록에서 계산 / 초록·전문 판독 = 요약 단계가 초록(본문을 읽은 논문은 본문)을 읽고 센 값 (해당 논문 키로 대조 가능)", 9)
    rpt.table(["주장", "편수", "기준", "해당 논문"],
              [[x["text"], x["count"], claim_basis(x, c.level_of), ", ".join(f"[{k}]" for k in x.get("keys", [])) or "—"] for x in claims],
              [6.0, 1.2, 1.9, 7.4], size=8)


def sec_appendix_reviews(rpt, c):
    rpt.table(["인용 키", "제목", "용도", "읽으면 좋은 이유"],
              [[r["key"] + (" ★먼저" if r.get("core") else ""), c.refs[r["key"]]["title"], r.get("use", ""), r["why_read"]]
               for r in c.s["reviews_guide"]], [2.6, 6.2, 2.2, 5.5], size=8.5)


def sec_appendix_adjacent(rpt, c):  # 주목할 인접 문헌 (사용자 결정 2026-10-16)
    rpt.para(ADJACENT_NOTE, 9)
    rows = adjacent_rows(c.s, c.refs)
    if not rows:
        rpt.para("해당 문헌 없음 — 기준(관련성·영향 점수 하한)을 넘는 논문이 없거나 요약에 목록이 없음", 9)
        return
    rpt.table(["인용 키", "제목", "연도·실린 곳", "피인용", "관련성", "알아 둘 이유"], rows, [2.2, 5.0, 2.8, 1.2, 1.2, 4.1], size=8.5)


def sec_appendix_data(rpt, c):
    rows = [[d["name"], d.get("kind") or "", d.get("access") or "접근 미확인", d.get("link") or "—", "\n".join(d.get("used_by", []))]
            for d in c.s.get("datasets_benchmarks", [])]
    if rows:
        rpt.para("접근 상태: 공개 접근 확인 = 링크 확인됨 / 논문 언급만 = 논문이 사용했다고만 밝힘 (바로 받을 수 있다는 뜻이 아님)", 9)
        rpt.table(["이름", "종류", "접근 상태", "링크", "사용 논문"], rows, [4.8, 2.4, 2.4, 3.6, 3.3], size=8.5)
    else:
        rpt.bullet("초록에서 확인된 데이터·벤치마크 없음")


def sec_appendix_terms(rpt, c):
    for t in c.s["core_technologies"]:
        rpt.bullet(f"{t['name']}: {t['description']}{cite(t)}")
    abbreviations = c.s.get("abbreviations") or {}
    if abbreviations:
        rpt.table(["약어", "정식 명칭"], [[a, f] for a, f in abbreviations.items()], [3.0, 13.5], size=9)


def sec_references(rpt, c):
    """APA 7판, 알파벳순, 내어쓰기. 본문 인용 키 [Tan2025] = 첫 저자 성 + 연도."""
    for key, r in sorted(c.refs.items(), key=apa_sort_key):
        p = rpt.doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(1.0)
        p.paragraph_format.first_line_indent = Cm(-1.0)
        p.paragraph_format.space_after = Pt(2)
        for text, italic in apa_reference(key, r):
            rpt.run(p, text, 9).italic = italic


SECTIONS = {  # id → (제목, 함수, 종류: main = 번호 / appendix = 부록 문자 / plain = 번호 없음)
    "summary": ("핵심 인사이트", sec_summary, "main"),
    "landscape": ("연구 접근과 문헌 분포", sec_landscape, "main"),
    "subtopics": ("분야별 대표 연구", sec_subtopics, "main"),
    "issues": ("후속 검토 쟁점", sec_issues, "main"),
    "conclusion": ("결론", sec_conclusion, "main"),
    "appendix_purpose": ("조사 목적과 범위", sec_purpose, "appendix"),
    "appendix_method": ("검색·선정 방법과 한계", sec_method, "appendix"),
    "appendix_next": ("다음 단계 — 전문 분석 대상", sec_appendix_next, "appendix"),
    # 예전 장 id (개인 설정에서 다시 켤 수 있음)
    "purpose": ("조사 목적과 범위", sec_purpose, "main"),
    "method": ("검색·선정 방법과 한계", sec_method, "main"),
    "appendix_excel": ("상세 자료 안내 (05_참고문헌.xlsx)", sec_appendix_excel, "appendix"),
    "appendix_methods": ("핵심논문 비교표", sec_appendix_methods, "appendix"),
    "appendix_excluded": ("선정 밖 문헌 점검표", sec_appendix_excluded, "appendix"),
    "appendix_counts": ("집계 근거표 — 본문 편수의 출처", sec_appendix_counts, "appendix"),
    "appendix_reviews": ("문헌조사논문 (핵심논문과 별도)", sec_appendix_reviews, "appendix"),
    "appendix_adjacent": ("주목할 인접 문헌 (핵심논문과 별도)", sec_appendix_adjacent, "appendix"),
    "appendix_data": ("데이터·벤치마크", sec_appendix_data, "appendix"),
    "appendix_terms": ("기법·용어", sec_appendix_terms, "appendix"),
    "references": ("참고문헌", sec_references, "plain"),
}


class Context:
    """장 함수들이 함께 쓰는 입력."""

    def __init__(self, run_dir: Path, variant: str):
        self.run_dir, self.vname = run_dir, lambda name: variant_name(name, variant)
        self.s = read_json(run_dir / self.vname("synthesis.json"))
        self.kw = self.s.get("keywords", [])  # 핵심 키워드 (없으면 예전 형식)
        scored = read_json(run_dir / "scored.json")
        self.meta, self.records = scored["meta"], scored["records"]
        self.refs = read_json(run_dir / "references.json")
        self.config = load_config()
        self.axis = map_axis(run_dir)  # 문헌지도·분포 가로축 (주제별)
        self.method_of = {m["key"]: m for m in self.s["method_comparison"]}
        self.review_line, self.level_of = review_info(run_dir, variant)
        self.number, self.figures = 0, 0

    def fig(self) -> int:
        self.figures += 1
        return self.figures


def build(run_dir: Path, variant: str) -> Path:
    c = Context(run_dir, variant)
    s, p = c.s, c.meta["prisma"]
    rpt = Report(s.get("abbreviations"), verbatim_strings(run_dir))
    rpt.tech = tech_keywords(s)  # 기술 핵심어 초록색 (사용자 피드백 2026-10-19)
    years = sorted({r["year"] for r in c.records})

    rpt.para("ALT 문헌조사 보고서", 12, True, NAVY, WD_ALIGN_PARAGRAPH.CENTER, mark=False)
    rpt.para(report_title(run_dir, s, c.refs), 18, True, NAVY, WD_ALIGN_PARAGRAPH.CENTER, mark=False)
    rpt.para(f"대상 출판 {years[0]}–{years[-1]} · 검색일 {c.meta['run_at'][:10]} · 후보 {p['screened']}편 → 핵심논문 {p['selected']}편 "
             f"+ 문헌조사논문 {p['reviews_listed']}편", 10, False, None, WD_ALIGN_PARAGRAPH.CENTER)
    rpt.para(c.review_line, 10, True, RED, WD_ALIGN_PARAGRAPH.CENTER)
    rpt.para("[ ] 안은 근거 문헌의 인용 키 (참고문헌·05_참고문헌.xlsx와 같음), [일반 개념]은 교과서 수준 지식",
             9, False, None, WD_ALIGN_PARAGRAPH.CENTER)

    appendix = 0
    for sid in sections("report", run_dir):
        title, fn, kind = SECTIONS[sid]
        if kind == "main":
            c.number += 1
            rpt.heading(f"{c.number}. {title}")
        elif kind == "appendix":
            appendix += 1
            rpt.heading(f"부록 {chr(64 + appendix)}. {title}")
        else:
            rpt.heading(title)
        rpt.lead((s.get("leads") or {}).get(sid, ""))
        fn(rpt, c)

    out = run_dir / c.vname("03_보고서.docx")
    rpt.doc.save(out)
    return out


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("run_dir", type=Path)
    add_variant(parser)
    args = parser.parse_args()
    print(f"[docx] {build(args.run_dir, args.variant).name}")


if __name__ == "__main__":
    main()
