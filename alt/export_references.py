"""참고문헌 내보내기 → 05_references.xlsx (보고서의 상세 표 전부 — 시트 목록은 EXCEL_SHEETS), 05_references.bib

    python alt/export_references.py outputs/<run>/
"""

from __future__ import annotations

import argparse
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from add_cocited import cocited_count
from common import add_variant, map_axis, read_json, review_info, tr, tr_later, tr_stored, use_language, variant_name
from report_facts import CATEGORY_LABEL, PDF_SOURCE_LABEL, adjacent_rows, claim_basis, excluded_comparison, excluded_rows, role_of

HEADER_FILL = PatternFill("solid", fgColor="D9D9D9")
STATUS_LABEL = {"selected": tr_later("핵심논문"), "review_listed": tr_later("문헌조사논문"), "review_not_listed": tr_later("리뷰(평가 탈락·목록 외)"),
                "below_relevance": tr_later("관련성 미달"), "below_threshold": tr_later("기준 점수 미달"), "over_limit": tr_later("상한 초과"),
                "excluded_retracted": tr_later("철회 제외"), "institution_listed": tr_later("기관·선급 문서"),
                "institution_not_listed": tr_later("기관·선급 문서(목록 외)"), "cocited_listed": tr_later("공통 인용 고전"),
                "cocited_not_listed": tr_later("공통 인용 고전(목록 외)"), "excluded_no_abstract": tr_later("초록 없음·인용 낮음 제외")}
# 보고서·발표의 '상세 자료' 안내에 그대로 쓰는 시트 목록 (상세 표는 보고서가 아니라 여기에 둔다)
# 이름·설명 모두 tr_later — 쓰는 곳(여기·build_docx·build_pptx·prepare_slides)에서 tr
EXCEL_SHEETS = [(tr_later("핵심논문"), tr_later("핵심논문 전체 — 점수·관련성 근거·서지·전문 확보")),
                (tr_later("방법 비교"), tr_later("핵심논문별 해결 문제·방법·입력·출력·초록상 검증 근거·전문 확인 필요")),
                (tr_later("문헌조사논문"), tr_later("문헌조사논문(리뷰) 서지·점수")), (tr_later("기관·선급 문서"), tr_later("선급 규칙·가이드, ISSC·ITTC·IMO 등 — 검색으로 찾은 것과 사용자 제공")), (tr_later("문헌조사논문 요지"), tr_later("리뷰가 정리한 내용과 쓰임")), (tr_later("주목할 인접 문헌"), tr_later("관련성은 낮지만 인용·학술지 점수가 높은 연구논문 (최대 5편, 본문 근거 아님)")),
                (tr_later("공통 인용 고전"), tr_later("핵심 후보들이 함께 인용하는데 검색에 안 걸린 분야 고전 (최대 10편, 공통 인용 많은 순 — 표시사항 '전문 확인 필요'는 초록 없이 제목으로 채점)")),
                (tr_later("선정 밖 목록"), tr_later("관련성은 통과했지만 총점 미달·상한 초과로 빠진 논문 (이유 칸)")), (tr_later("핵심·선정 밖 비교"), tr_later("핵심논문과 선정 밖 논문의 연도·종류·문헌지도 가로축 분포")),
                (tr_later("집계 근거"), tr_later("보고서 편수 주장마다 기준(자동 집계/초록·전문 판독)과 해당 논문 키")),
                (tr_later("데이터·벤치마크"), tr_later("논문이 쓴 데이터·모델과 접근 상태")), (tr_later("기법·약어"), tr_later("AI 기법 설명과 약어 정식 명칭")),
                (tr_later("전체 후보"), tr_later("검색 후보 전체와 제외 사유"))]


def pct(value):
    return None if value is None else round(value * 100)


def paper_row(rank, key, r, ref=None, level="") -> list:
    parts = r["score_parts"]
    ref = ref or {}
    verified = (f"{r['relevance_original']}→{r['verifier_relevance']}" if r.get("verified")
                and r.get("relevance_original") != r.get("verifier_relevance") else (tr("일치") if r.get("verified") else ""))
    return [rank, key, "★" * r["stars"], round(r["total_score"], 3), r["relevance"], r["relevance_reason"],
            pct(parts["citation"]), pct(parts["venue"]), tr(CATEGORY_LABEL[r["category"]]), r["title"],
            ", ".join(r["authors"][:3]) + (tr(" 외") if len(r["authors"]) > 3 else ""), r["year"], r["venue"],
            r["doi"], ref.get("oa_pdf_url") or ref.get("oa_landing_url") or r.get("oa_url"),
            "✓" if ref.get("pdf_file") else "", tr(PDF_SOURCE_LABEL.get(ref.get("pdf_status"), "")), tr(level), ref.get("pdf_file"),
            tr(STATUS_LABEL[r["status"]]), verified, " / ".join(tr_stored(f) for f in r["flags"]), tr(r.get("update") or "")]


PAPER_HEADER = [tr_later("순위"), tr_later("인용 키"), "★", tr_later("총점"), tr_later("관련성(0–5)"), tr_later("관련성 근거"),
                tr_later("인용 백분위"), tr_later("학술지 백분위"), tr_later("종류"), tr_later("제목"), tr_later("저자"), tr_later("연도"),
                tr_later("학술지·학회"), "DOI", tr_later("무료 공개본 링크"), tr_later("전문 저장"), tr_later("수집 경로"), tr_later("검토 수준"),
                tr_later("저장된 파일"), tr_later("상태"), tr_later("교차 검증(만든 AI→검증 AI)"), tr_later("표시사항"), tr_later("갱신")]
WIDTHS = [6, 14, 8, 7, 9, 50, 9, 9, 8, 60, 25, 6, 28, 30, 30, 8, 12, 9, 22, 12, 14, 40, 10]
LINK_COLUMNS = (14, 15, 19)


def write_sheet(ws, header, rows, widths, link_cols=()) -> None:
    ws.append([tr(h) for h in header])  # 머리글은 tr_later 원문·config 표시 이름 → 여기서 결과물 언어로
    for cell in ws[1]:
        cell.font = Font(bold=True, name="맑은 고딕")
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for row in rows:
        ws.append(row)
    for i, width in enumerate(widths, 1):
        ws.column_dimensions[ws.cell(1, i).column_letter].width = width
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.font = Font(name="맑은 고딕", size=10)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            if cell.column in link_cols and isinstance(cell.value, str) and cell.value.startswith(("http", "papers/")):
                cell.hyperlink = cell.value
                cell.font = Font(name="맑은 고딕", size=10, color="0563C1", underline="single")
    ws.freeze_panes = "C2"


def bibtex(ref: dict) -> str:
    kind = {"conference": "inproceedings", "institution": "techreport"}.get(ref["category"], "article")
    venue_field = {"inproceedings": "booktitle", "techreport": "institution"}.get(kind, "journal")
    fields = {"title": ref["title"], "author": " and ".join(ref["authors"]), "year": ref["year"],
              venue_field: ref["venue"], "number": ref.get("doc_number"), "doi": (ref["doi"] or "").replace("https://doi.org/", "") or None}
    body = ",\n".join(f"  {k} = {{{v}}}" for k, v in fields.items() if v)
    return f"@{kind}{{{ref['key']},\n{body}\n}}\n"


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("run_dir", type=Path)
    add_variant(parser)
    args = parser.parse_args()
    use_language(args.run_dir)

    scored = read_json(args.run_dir / "scored.json")
    references = read_json(args.run_dir / "references.json")
    synthesis = read_json(args.run_dir / variant_name("synthesis.json", args.variant))
    key_of = {ref["openalex_id"]: key for key, ref in references.items()}
    records = scored["records"]
    _, level_of = review_info(args.run_dir, args.variant)

    wb = Workbook()
    sheets = [("핵심논문", lambda r: r["status"] == "selected"), ("문헌조사논문", lambda r: r["status"] == "review_listed"),
              ("공통 인용 고전", lambda r: r["status"] == "cocited_listed"), ("전체 후보", lambda r: True)]
    for i, (name, keep) in enumerate(sheets):
        chosen = [r for r in records if keep(r)]
        if name == "공통 인용 고전":
            chosen.sort(key=lambda r: -cocited_count(r))
        ws = wb.active if i == 0 else wb.create_sheet()
        ws.title = tr(name)
        rows = []
        for n, r in enumerate(chosen, 1):
            key = key_of.get(r["openalex_id"], "")
            rows.append(paper_row(n, key, r, references.get(key), level_of.get(key, "")))
        write_sheet(ws, PAPER_HEADER, rows, WIDTHS, LINK_COLUMNS)

    ws = wb.create_sheet(tr("데이터·벤치마크"))
    rows = [[d["name"], tr(d.get("kind") or ""), tr(d.get("access") or "접근 미확인"), d.get("link") or "", ", ".join(d.get("used_by", []))]
            for d in synthesis.get("datasets_benchmarks", [])]
    write_sheet(ws, [tr_later("이름"), tr_later("종류"), tr_later("접근 상태"), tr_later("링크"), tr_later("사용한 논문(인용 키)")],
                rows, [40, 16, 16, 40, 40], (4,))

    write_sheet(wb.create_sheet(tr("기관·선급 문서")),
                [tr_later("인용 키"), tr_later("기관·발행처"), tr_later("제목"), tr_later("연도"), tr_later("문서 번호"), tr_later("출처"),
                 tr_later("관련성"), tr_later("DOI·파일")],
                [[k, ref["venue"], ref["title"], ref["year"], ref.get("doc_number") or "", tr("사용자 제공") if ref.get("user_provided") else tr("검색"),
                  ref.get("relevance") if ref.get("relevance") is not None else "", ref.get("doi") or ref.get("source_file") or ""]
                 for k, ref in references.items() if ref["category"] == "institution"], [14, 28, 60, 7, 18, 12, 8, 36])

    # 보고서에서 옮겨 온 상세 표 (보고서에는 시트 안내만 남김)
    axis = map_axis(args.run_dir)
    role = role_of(synthesis)
    subtopic_of = {k: t["name"] for t in synthesis.get("subtopics", []) for k in t["papers"]}
    write_sheet(wb.create_sheet(tr("방법 비교")),
                [tr_later("인용 키"), tr_later("분야"), tr_later("역할"), tr_later("관련성"), tr_later("검토 수준"), tr_later("해결 문제"),
                 tr_later("방법"), tr_later("입력"), tr_later("출력"), tr_later("검증 근거"), tr_later("학습·검증 자료(판독)"),
                 tr_later("전문 확인 필요"), tr_later("구조물")],
                [[m["key"], subtopic_of.get(m["key"], ""), role.get(m["key"], ""), references[m["key"]]["relevance"],
                  tr(level_of.get(m["key"], "")), m.get("problem", ""), m["method"], m["input"], m["output"], m["validation"],
                  ", ".join(tr(x) for x in (m.get("data_source") or ["미기재"])), m.get("to_verify", ""), m["structure"]]
                 for m in synthesis.get("method_comparison", [])],
                [14, 22, 7, 8, 9, 34, 40, 34, 34, 40, 16, 34, 20])
    write_sheet(wb.create_sheet(tr("문헌조사논문 요지")),
                [tr_later("인용 키"), tr_later("제목"), tr_later("쓰임"), tr_later("핵심"), tr_later("정리한 내용")],
                [[r["key"], references[r["key"]]["title"], tr(r.get("use", "")), "★" if r.get("core") else "", r["why_read"]]
                 for r in synthesis.get("reviews_guide", [])], [16, 60, 14, 9, 60])
    write_sheet(wb.create_sheet(tr("주목할 인접 문헌")),
                [tr_later("인용 키"), tr_later("제목"), tr_later("연도·실린 곳"), tr_later("피인용"), tr_later("관련성"),
                 tr_later("알아 둘 이유"), tr_later("영향 점수")],
                [row + [references[row[0]].get("impact_score")] for row in adjacent_rows(synthesis, references)], [14, 60, 30, 8, 8, 60, 9])
    if scored["meta"].get("update"):  # 누적 갱신: 이전 핵심논문 중 이번에 빠진 논문 (사용자 지정 방식)
        dropped = sorted((r for r in records if r.get("update") == "제외됨(이전 선정)"), key=lambda r: -r["total_score"])
        write_sheet(wb.create_sheet(tr("이전 선정 제외")),
                    [tr_later("제목"), tr_later("연도"), tr_later("지금 상태"), tr_later("관련성"), tr_later("총점"), tr_later("이전 결과")],
                    [[r["title"], r["year"], tr(STATUS_LABEL.get(r["status"], r["status"])), r["relevance"], round(r["total_score"], 3),
                      scored["meta"]["update"]["base"]] for r in dropped], [70, 7, 22, 8, 8, 16])
    write_sheet(wb.create_sheet(tr("선정 밖 목록")),
                [tr_later("순위"), tr_later("제목"), tr_later("연도"), tr_later("종류"), tr_later("인용 수"), tr_later("관련성"),
                 tr_later("총점"), axis["name"], tr_later("이유")],
                excluded_rows(records, axis), [6, 70, 7, 10, 8, 8, 8, 18, 10])
    write_sheet(wb.create_sheet(tr("핵심·선정 밖 비교")), [tr_later("항목"), tr_later("핵심논문"), tr_later("선정 밖")],
                excluded_comparison(records, axis), [28, 10, 10])
    write_sheet(wb.create_sheet(tr("집계 근거")), [tr_later("주장"), tr_later("편수"), tr_later("기준"), tr_later("해당 논문(인용 키)")],
                [[x["text"], x["count"], claim_basis(x, level_of), ", ".join(x.get("keys", []))] for x in synthesis.get("count_claims", [])],
                [60, 7, 10, 70])
    write_sheet(wb.create_sheet(tr("기법·약어")), [tr_later("구분"), tr_later("이름"), tr_later("설명"), tr_later("근거 인용 키")],
                [[tr("기법"), t["name"], t["description"], ", ".join(t.get("cites", []))] for t in synthesis.get("core_technologies", [])]
                + [[tr("약어"), a, f, ""] for a, f in (synthesis.get("abbreviations") or {}).items()], [8, 30, 70, 30])
    if scored["meta"].get("scoring_overrides"):
        write_sheet(wb.create_sheet(tr("기본값과 다른 기준")), [tr_later("항목"), tr_later("기본값"), tr_later("사용한 값"), tr_later("출처")],
                    [[o["key"], str(o["default"]), str(o["value"]), tr(o["source"])] for o in scored["meta"]["scoring_overrides"]], [30, 20, 20, 14])
    order = [tr(name) for name, _ in EXCEL_SHEETS]
    wb._sheets.sort(key=lambda ws: order.index(ws.title) if ws.title in order else len(order))

    xlsx = args.run_dir / variant_name("05_references.xlsx", args.variant)
    wb.save(xlsx)
    bib = args.run_dir / variant_name("05_references.bib", args.variant)
    bib.write_text("\n".join(bibtex(r) for r in references.values()), encoding="utf-8")
    print(f"[references] {xlsx.name} (핵심논문 {sum(r['status'] == 'selected' for r in records)}, "
          f"문헌조사논문 {sum(r['status'] == 'review_listed' for r in records)}, 전체 {len(records)}, "
          f"데이터 {len(rows)}), {bib.name} {len(references)}건")


if __name__ == "__main__":
    main()
